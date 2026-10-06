"""Turns an uploaded CSV or Excel file into customers, orders and order items.

The file has one row per order line (one product in one order):

    required: order_id, order_date, customer_id, product_code, quantity, unit_price
    optional: customer_name, customer_email, product_name

Rules:
- A row with a problem is rejected and listed in the error report. The other
  rows still import.
- All lines of one order must have the same order_date and customer_id.
- An order whose order_id was imported before is skipped, so uploading the
  same file twice does not double the revenue.
- A customer seen before (same customer_id) is reused, and a new name or
  email from the file updates it.
"""

import csv
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import delete, exists, func, select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.models import Customer, ImportJob, ImportStatus, Order, OrderItem

REQUIRED_COLUMNS = (
    "order_id",
    "order_date",
    "customer_id",
    "product_code",
    "quantity",
    "unit_price",
)
OPTIONAL_COLUMNS = ("customer_name", "customer_email", "product_name")
SUPPORTED_SUFFIXES = (".csv", ".xlsx")

# The error report keeps the first problems only. The count of rejected rows is always exact.
MAX_REPORTED_ERRORS = 1000
BATCH_SIZE = 1000
LARGEST_AMOUNT = Decimal("1e10")


class FileError(Exception):
    """The file as a whole can't be imported (wrong type, missing columns...)."""


@dataclass(slots=True)
class Line:
    """One valid row of the file."""

    row: int
    order_id: str
    ordered_at: datetime
    customer_id: str
    customer_name: str | None
    customer_email: str | None
    product_code: str
    product_name: str | None
    quantity: int
    unit_price: Decimal


@dataclass
class ImportResult:
    rows_imported: int = 0
    rows_rejected: int = 0
    rows_skipped: int = 0
    orders_imported: int = 0


class ErrorReport:
    def __init__(self) -> None:
        self.entries: list[dict] = []
        self.rejected_rows = 0

    def add(self, row: int, errors: list[str]) -> None:
        self.rejected_rows += 1
        if len(self.entries) < MAX_REPORTED_ERRORS:
            self.entries.append({"row": row, "errors": errors})


# --- Reading the file -------------------------------------------------------


def read_rows(path: Path) -> Iterator[tuple[int, dict[str, object]]]:
    """Yield (row number as shown in a spreadsheet, {column: value}) for each non-empty row."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        yield from _read_csv(path)
    elif suffix == ".xlsx":
        yield from _read_xlsx(path)
    else:
        raise FileError("Upload a .csv or .xlsx file.")


def _read_csv(path: Path) -> Iterator[tuple[int, dict[str, object]]]:
    try:
        # utf-8-sig also accepts the invisible marker Excel puts at the start of "CSV UTF-8" files.
        with open(path, newline="", encoding="utf-8-sig") as file:
            reader = csv.reader(file)
            header = _clean_header(next(reader, []))
            for row_number, cells in enumerate(reader, start=2):
                if any(cell.strip() for cell in cells):
                    yield row_number, dict(zip(header, cells, strict=False))
    except UnicodeDecodeError:
        raise FileError(
            'The file isn\'t UTF-8 text. In Excel, save it as "CSV UTF-8" and upload it again.'
        ) from None
    except csv.Error as exc:
        raise FileError(f"The CSV file is malformed: {exc}.") from None


def _read_xlsx(path: Path) -> Iterator[tuple[int, dict[str, object]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError):
        raise FileError(
            "The Excel file couldn't be opened. Check it's a valid .xlsx file."
        ) from None
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)  # first sheet only
        header = _clean_header(next(rows, ()))
        for row_number, cells in enumerate(rows, start=2):
            if any(cell is not None and str(cell).strip() for cell in cells):
                yield row_number, dict(zip(header, cells, strict=False))
    finally:
        workbook.close()


def _clean_header(cells: Iterable[object]) -> list[str]:
    """Lowercase the column names and accept spaces for underscores ("Order ID" -> "order_id")."""
    header = [str(cell or "").strip().lower().replace(" ", "_") for cell in cells]
    missing = [column for column in REQUIRED_COLUMNS if column not in header]
    if missing:
        raise FileError(f"The file is missing required columns: {', '.join(missing)}.")
    return header


# --- Checking each row ------------------------------------------------------


def _as_text(value: object) -> str:
    # Excel stores numbers like customer IDs as floats: 17850.0 should read as "17850".
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return "" if value is None else str(value).strip()


def parse_line(row_number: int, raw: dict[str, object]) -> Line | list[str]:
    """Return a valid Line, or the list of problems with this row."""
    errors: list[str] = []

    def text(column: str, max_length: int, required: bool) -> str | None:
        value = _as_text(raw.get(column))
        if not value:
            if required:
                errors.append(f"{column} is empty")
            return None
        if len(value) > max_length:
            errors.append(f"{column} is longer than {max_length} characters")
            return None
        return value

    order_id = text("order_id", 100, required=True)
    customer_id = text("customer_id", 100, required=True)
    product_code = text("product_code", 64, required=True)
    customer_name = text("customer_name", 200, required=False)
    customer_email = text("customer_email", 320, required=False)
    product_name = text("product_name", 1000, required=False)
    if customer_email and "@" not in customer_email:
        errors.append("customer_email doesn't look like an email address")

    ordered_at = _parse_date(raw.get("order_date"), errors)
    quantity = _parse_quantity(raw.get("quantity"), errors)
    unit_price = _parse_price(raw.get("unit_price"), errors)

    if errors:
        return errors
    return Line(
        row=row_number,
        order_id=order_id,
        ordered_at=ordered_at,
        customer_id=customer_id,
        customer_name=customer_name,
        customer_email=customer_email,
        product_code=product_code,
        product_name=product_name,
        quantity=quantity,
        unit_price=unit_price,
    )


def _parse_date(value: object, errors: list[str]) -> datetime | None:
    """Accept Excel dates and ISO text such as 2024-03-31 or 2024-03-31 14:05.
    Times without a timezone are treated as UTC."""
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime(value.year, value.month, value.day)
    else:
        text = _as_text(value)
        if not text:
            errors.append("order_date is empty")
            return None
        try:
            parsed = datetime.fromisoformat(text)
        except ValueError:
            errors.append(f'order_date "{text}" isn\'t a date like 2024-03-31 or 2024-03-31 14:05')
            return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)


def _parse_quantity(value: object, errors: list[str]) -> int | None:
    text = _as_text(value)
    try:
        quantity = Decimal(text)
    except InvalidOperation:
        errors.append(f'quantity "{text}" isn\'t a number' if text else "quantity is empty")
        return None
    if (
        not quantity.is_finite()
        or quantity != quantity.to_integral_value()
        or not 0 < quantity <= 1_000_000_000
    ):
        errors.append(f'quantity "{text}" must be a positive whole number')
        return None
    return int(quantity)


def _parse_price(value: object, errors: list[str]) -> Decimal | None:
    # str() first so a float from Excel like 2.55 doesn't become 2.54999...
    text = _as_text(value)
    try:
        price = Decimal(text)
    except InvalidOperation:
        errors.append(f'unit_price "{text}" isn\'t a number' if text else "unit_price is empty")
        return None
    if not price.is_finite() or price < 0 or price >= LARGEST_AMOUNT:
        errors.append(f'unit_price "{text}" must be zero or more')
        return None
    return price.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)


# --- Saving -----------------------------------------------------------------


def import_file(
    db: Session, organization_id: int, import_job_id: int, path: Path
) -> tuple[ImportResult, ErrorReport]:
    """Validate the file and save its new orders, tagged with the import. Does not commit."""
    report = ErrorReport()
    result = ImportResult()

    # 1. Check every row and group the valid ones by order.
    lines_by_order: dict[str, list[Line]] = {}
    for row_number, raw in read_rows(path):
        parsed = parse_line(row_number, raw)
        if isinstance(parsed, list):
            report.add(row_number, parsed)
        else:
            lines_by_order.setdefault(parsed.order_id, []).append(parsed)

    # 2. Every line of an order must agree on date and customer with the order's first line.
    for order_id, lines in lines_by_order.items():
        first = lines[0]
        consistent = []
        for line in lines:
            if line.ordered_at == first.ordered_at and line.customer_id == first.customer_id:
                consistent.append(line)
            else:
                report.add(
                    line.row,
                    [
                        f"order {order_id} has a different order_date or customer_id "
                        f"on row {first.row}"
                    ],
                )
        lines_by_order[order_id] = consistent

    # 3. Skip orders that were imported before.
    for batch in _batches(list(lines_by_order)):
        already_imported = db.scalars(
            select(Order.external_id).where(
                Order.organization_id == organization_id, Order.external_id.in_(batch)
            )
        )
        for order_id in already_imported:
            result.rows_skipped += len(lines_by_order.pop(order_id))

    result.rows_rejected = report.rejected_rows
    if not lines_by_order:
        return result, report

    # 4. Create new customers and update existing ones. Later rows win for name and email.
    customers: dict[str, dict] = {}
    for lines in lines_by_order.values():
        for line in lines:
            customer = customers.setdefault(line.customer_id, {"name": None, "email": None})
            customer["name"] = line.customer_name or customer["name"]
            customer["email"] = line.customer_email or customer["email"]

    customer_ids: dict[str, int] = {}
    for batch in _batches(list(customers.items())):
        statement = insert(Customer).values(
            [
                {
                    "organization_id": organization_id,
                    "external_id": external_id,
                    "name": details["name"],
                    "email": details["email"],
                }
                for external_id, details in batch
            ]
        )
        # "Upsert": if the customer exists, keep the old name/email unless the file has new ones.
        statement = statement.on_conflict_do_update(
            index_elements=["organization_id", "external_id"],
            set_={
                "name": func.coalesce(statement.excluded.name, Customer.name),
                "email": func.coalesce(statement.excluded.email, Customer.email),
            },
        ).returning(Customer.external_id, Customer.id)
        customer_ids.update(dict(db.execute(statement).tuples().all()))

    # 5. Insert orders, then their items.
    order_ids: dict[str, int] = {}
    for batch in _batches(list(lines_by_order.items())):
        statement = (
            insert(Order)
            .values(
                [
                    {
                        "organization_id": organization_id,
                        "import_job_id": import_job_id,
                        "customer_id": customer_ids[lines[0].customer_id],
                        "external_id": order_id,
                        "ordered_at": lines[0].ordered_at,
                        "total_amount": order_total(lines),
                    }
                    for order_id, lines in batch
                ]
            )
            .returning(Order.external_id, Order.id)
        )
        order_ids.update(dict(db.execute(statement).tuples().all()))

    all_lines = [line for lines in lines_by_order.values() for line in lines]
    for batch in _batches(all_lines):
        db.execute(
            insert(OrderItem).values(
                [
                    {
                        "organization_id": organization_id,
                        "order_id": order_ids[line.order_id],
                        "product_code": line.product_code,
                        "product_name": line.product_name,
                        "quantity": line.quantity,
                        "unit_price": line.unit_price,
                    }
                    for line in batch
                ]
            )
        )

    result.rows_imported = len(all_lines)
    result.orders_imported = len(order_ids)
    return result, report


def order_total(lines: list[Line]) -> Decimal:
    total = sum((line.quantity * line.unit_price for line in lines), Decimal(0))
    return total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _batches(items: list) -> Iterator[list]:
    for start in range(0, len(items), BATCH_SIZE):
        yield items[start : start + BATCH_SIZE]


# --- Running a job ----------------------------------------------------------


def stored_file_path(upload_dir: str, job: ImportJob) -> Path:
    """Where the uploaded file for a job is kept until it is processed."""
    return Path(upload_dir) / f"{job.id}{Path(job.filename).suffix.lower()}"


def process_import(db: Session, job_id: int, upload_dir: str) -> None:
    """Run one import job from start to finish and record the outcome on the job.

    Imported rows and the recalculated metrics are saved in one transaction:
    if anything fails, no rows are saved and the job is marked failed.
    """
    from app.metrics import recalculate_metrics

    job = db.get(ImportJob, job_id)
    if job is None or job.status != ImportStatus.QUEUED:
        return
    job.status = ImportStatus.RUNNING
    job.started_at = datetime.now(UTC)
    db.commit()

    path = stored_file_path(upload_dir, job)
    try:
        result, report = import_file(db, job.organization_id, job.id, path)
        recalculate_metrics(db, job.organization_id)
    except FileError as exc:
        db.rollback()
        job.status = ImportStatus.FAILED
        job.failure_reason = str(exc)
    except Exception:
        db.rollback()
        job.status = ImportStatus.FAILED
        job.failure_reason = "Something went wrong while importing. No rows were saved."
        job.finished_at = datetime.now(UTC)
        db.commit()
        raise  # so the worker log shows the full error
    else:
        job.status = ImportStatus.SUCCEEDED
        job.rows_imported = result.rows_imported
        job.rows_rejected = result.rows_rejected
        job.rows_skipped = result.rows_skipped
        job.error_report = report.entries or None
    finally:
        path.unlink(missing_ok=True)

    job.finished_at = datetime.now(UTC)
    db.commit()


# --- Deleting an import -----------------------------------------------------


def delete_import(db: Session, job: ImportJob) -> None:
    """Remove an import's orders, any customers left with no orders, and the job
    itself, then recalculate the metrics. Does not commit.

    Name or email changes the import made to customers who still have other
    orders are not undone: the database keeps only the latest values.
    """
    from app.metrics import recalculate_metrics

    org = job.organization_id
    orders_of_job = select(Order.id).where(
        Order.organization_id == org, Order.import_job_id == job.id
    )
    db.execute(
        delete(OrderItem).where(
            OrderItem.organization_id == org, OrderItem.order_id.in_(orders_of_job)
        )
    )
    db.execute(delete(Order).where(Order.organization_id == org, Order.import_job_id == job.id))

    # Rebuilds the daily tables and drops metrics of customers with no orders left.
    recalculate_metrics(db, org)

    has_orders = exists().where(Order.organization_id == org, Order.customer_id == Customer.id)
    db.execute(delete(Customer).where(Customer.organization_id == org, ~has_orders))
    db.delete(job)


def process_delete(db: Session, job_id: int) -> None:
    """Worker side of deleting an import that has data. All or nothing."""
    job = db.get(ImportJob, job_id)
    if job is None or job.status != ImportStatus.DELETING:
        return
    try:
        delete_import(db, job)
        db.commit()
    except Exception:
        db.rollback()
        job = db.get(ImportJob, job_id)
        job.status = ImportStatus.SUCCEEDED  # nothing was removed; let the user try again
        db.commit()
        raise
