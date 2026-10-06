"""Load a public dataset as a demo organization.

Dataset: "Online Retail" from the UCI Machine Learning Repository
(https://archive.ics.uci.edu/dataset/352/online+retail), licensed CC BY 4.0.
About 540,000 order lines from a UK online gift shop, Dec 2010 to Dec 2011.

The spreadsheet is converted to this app's import format and loaded through
the same import code as a normal upload. Rows the app doesn't take are
dropped first: cancellations (invoice numbers starting with "C"), lines
without a customer, and lines with zero or negative quantity or price. Some
invoices have lines stamped a minute apart; every line gets the invoice's
first timestamp, since the importer requires one date per order.

Run it with:
    docker compose exec api python -m scripts.seed
    docker compose exec api python -m scripts.seed --email you@example.com --password ...
"""

import argparse
import csv
import secrets
import sys
import urllib.request
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.config import settings
from app.db import SessionLocal
from app.importer import process_import, stored_file_path
from app.models import ImportJob, Organization, Role, User
from app.security import hash_password

DATASET_URL = "https://archive.ics.uci.edu/static/public/352/online+retail.zip"
CACHE_DIR = Path("/data/seed")
ORGANIZATION_NAME = "Online Retail (demo)"


def download_dataset() -> Path:
    """Download and unzip the dataset once; later runs reuse the copy."""
    workbook = CACHE_DIR / "Online Retail.xlsx"
    if not workbook.exists():
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        archive = CACHE_DIR / "online_retail.zip"
        print(f"Downloading {DATASET_URL} (about 23 MB)...")
        urllib.request.urlretrieve(DATASET_URL, archive)
        with zipfile.ZipFile(archive) as zipped:
            zipped.extractall(CACHE_DIR)
    return workbook


def convert(workbook_path: Path, csv_path: Path) -> tuple[int, int]:
    """Write the rows we keep as an import CSV. Returns (rows kept, rows dropped)."""
    print("Converting the spreadsheet (this takes about a minute)...")
    kept = dropped = 0
    invoice_dates = {}  # invoice -> timestamp of its first line
    workbook = load_workbook(workbook_path, read_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    next(rows)  # header: InvoiceNo, StockCode, Description, Quantity, InvoiceDate, ...
    with open(csv_path, "w", newline="", encoding="utf-8") as out:
        writer = csv.writer(out)
        writer.writerow(
            [
                "order_id",
                "order_date",
                "customer_id",
                "product_code",
                "product_name",
                "quantity",
                "unit_price",
            ]
        )
        for invoice, code, description, quantity, invoiced_at, price, customer, _ in rows:
            if (
                customer is None
                or str(invoice).startswith("C")
                or not quantity
                or quantity <= 0
                or not price
                or price <= 0
            ):
                dropped += 1
                continue
            writer.writerow(
                [
                    invoice,
                    invoice_dates.setdefault(invoice, invoiced_at).isoformat(),
                    int(customer),
                    code,
                    (description or "").strip(),
                    quantity,
                    price,
                ]
            )
            kept += 1
    workbook.close()
    return kept, dropped


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--email", default="demo@example.com", help="admin login to create")
    parser.add_argument("--password", help="admin password (a random one is made if omitted)")
    args = parser.parse_args()
    password = args.password or secrets.token_urlsafe(12)
    email = args.email.lower()

    with SessionLocal() as db:
        if db.scalar(select(User).where(User.email == email)):
            sys.exit(f"{email} already exists. Pick another --email or run: docker compose down -v")

        workbook = download_dataset()

        organization = Organization(name=ORGANIZATION_NAME, currency="GBP")
        admin = User(
            organization=organization,
            email=email,
            password_hash=hash_password(password),
            role=Role.ADMIN,
        )
        db.add(admin)
        db.flush()
        job = ImportJob(
            organization_id=organization.id,
            created_by_user_id=admin.id,
            filename="online_retail.csv",
        )
        db.add(job)
        db.flush()

        csv_path = stored_file_path(settings.upload_dir, job)
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        kept, dropped = convert(workbook, csv_path)
        print(f"Kept {kept:,} rows, dropped {dropped:,} (cancellations, no customer, no amount).")
        db.commit()

        print("Importing and calculating metrics...")
        process_import(db, job.id, settings.upload_dir)
        db.refresh(job)

    print()
    print(
        f"Import {job.status}: {job.rows_imported:,} rows imported, "
        f"{job.rows_rejected:,} rejected, {job.rows_skipped:,} skipped."
    )
    if job.failure_reason:
        print(f"Reason: {job.failure_reason}")
    print()
    print("Log in at http://localhost:3000 with")
    print(f"  email:    {email}")
    print(f"  password: {password}")


if __name__ == "__main__":
    main()
